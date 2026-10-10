"""Build the GitHub Pages gallery: thumbs/ + gallery_data.json.

Inputs (all in the repo): Books_Shopify_Import.csv, Books_Shopify_Review.xlsx,
photos_report.csv and the photos/ folder. Run from the repo root:
    python scripts/build_site.py
"""
import csv
import json
import os
import re

from openpyxl import load_workbook
from PIL import Image

THUMB_H = 360  # 2x the 180px card image height


def thumb_path(photo):
    return 'thumbs/' + os.path.splitext(photo[len('photos/'):])[0] + '.jpg'


def make_thumb(src, dst):
    if os.path.exists(dst):
        return
    os.makedirs(os.path.dirname(dst), exist_ok=True)
    with Image.open(src) as im:
        im.seek(0)
        im = im.convert('RGBA')
        bg = Image.new('RGB', im.size, (255, 255, 255))
        bg.paste(im, mask=im.split()[3])
        bg.thumbnail((THUMB_H * 2, THUMB_H))
        bg.save(dst, 'JPEG', quality=80, optimize=True, progressive=True)


def main():
    rows = list(csv.DictReader(open('Books_Shopify_Import.csv', encoding='utf-8')))
    report = {}
    for r in csv.DictReader(open('photos_report.csv', encoding='utf-8')):
        if r['status'] == 'ok':
            report[(r['handle'], int(r['position']))] = r['file']

    ws = load_workbook('Books_Shopify_Review.xlsx', read_only=True)['All products']
    it = ws.iter_rows(values_only=True)
    head = next(it)
    review = {str(row[head.index('Barcode')]): dict(zip(head, row)) for row in it}

    products, current = [], None
    for r in rows:
        if r['Title']:
            bc = r['Variant Barcodes']
            rv = review.get(bc, {})
            current = {
                'title': r['Title'],
                'sheet_title': rv.get('Sheet title') or '',
                'barcode': bc,
                'sku': r['SKU'],
                'price': r['Price'],
                'type': r['Type'],
                'author': rv.get('Author') or '',
                'vendor': r['Vendor'],
                'language': rv.get('Language') or '',
                'status': r['Status'],
                'match': rv.get('Match result') or '',
                'photo_source': rv.get('Photo source') or '',
                'attention': rv.get('Needs attention') or '',
                'description': r['Description'],
                'photos': [],
            }
            products.append(current)
        if r['Product image URL']:
            f = report.get((r['URL handle'], int(r['Image position'])))
            if f:
                t = thumb_path(f)
                make_thumb(f, t)
                current['photos'].append({'full': f, 'thumb': t})

    with open('gallery_data.json', 'w', encoding='utf-8') as fh:
        json.dump(products, fh, ensure_ascii=False, separators=(',', ':'))
    n_ph = sum(len(p['photos']) for p in products)
    print(f'{len(products)} products, {n_ph} photos, {sum(1 for p in products if not p["photos"])} without photos')


if __name__ == '__main__':
    main()
